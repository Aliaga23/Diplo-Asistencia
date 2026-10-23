import io
from uuid import UUID

from fastapi import APIRouter, HTTPException, status
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.core.deps import DB, AdminUser
from app.core.timeutils import now_local, today_local
from app.models import AttendanceStatus, ClassSession, Course, Enrollment, Student, User
from app.schemas import DashboardStats, SessionPoint
from app.services.attendance import enrolled_count, is_closed, session_out

router = APIRouter(tags=["Panel y reportes"])

STATUS_LABEL = {
    AttendanceStatus.present: "P",
    AttendanceStatus.late: "T",
    AttendanceStatus.absent: "F",
    AttendanceStatus.excused: "J",
}


def _course_or_404(db, course_id: UUID | None) -> Course:
    course = (
        db.get(Course, course_id)
        if course_id
        else db.scalar(select(Course).order_by(Course.created_at).limit(1))
    )
    if course is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "No hay cursos registrados")
    return course


@router.get("/dashboard/stats", response_model=DashboardStats, summary="Indicadores del panel")
def dashboard(db: DB, _: AdminUser, course_id: UUID | None = None) -> DashboardStats:
    course = _course_or_404(db, course_id)
    today = today_local()
    enrolled = enrolled_count(db, course.id)
    sessions = db.scalars(
        select(ClassSession)
        .options(selectinload(ClassSession.attendance_records))
        .where(ClassSession.course_id == course.id)
        .order_by(ClassSession.session_date, ClassSession.start_time)
    ).all()
    done = [s for s in sessions if is_closed(s)]
    outs = [session_out(db, s, enrolled) for s in done]
    attended = sum(o.present + o.late for o in outs)
    possible = enrolled * len(done)
    today_s = next((s for s in sessions if s.session_date == today), None)
    today_o = session_out(db, today_s, enrolled) if today_s else None
    series = [
        SessionPoint(
            session_id=o.id,
            label=o.session_date.strftime("%d/%m"),
            session_date=o.session_date,
            present=o.present,
            late=o.late,
            absent=o.absent,
            excused=o.excused,
            rate=round((o.present + o.late) / enrolled * 100, 1) if enrolled else 0.0,
        )
        for o in outs[-14:]
    ]
    return DashboardStats(
        students_total=enrolled,
        sessions_total=len(sessions),
        sessions_done=len(done),
        attendance_rate=round(attended / possible * 100, 1) if possible else 0.0,
        today_session=today_o,
        today_present=today_o.present if today_o else 0,
        today_late=today_o.late if today_o else 0,
        today_pending=(
            max(enrolled - today_o.present - today_o.late - today_o.excused, 0) if today_o else 0
        ),
        series=series,
        server_now=now_local(),
    )


@router.get("/reports/attendance.xlsx", summary="Exportar asistencia a Excel")
def export_attendance(db: DB, _: AdminUser, course_id: UUID | None = None) -> Response:
    course = _course_or_404(db, course_id)
    sessions = db.scalars(
        select(ClassSession)
        .options(selectinload(ClassSession.attendance_records))
        .where(ClassSession.course_id == course.id)
        .order_by(ClassSession.session_date, ClassSession.start_time)
    ).all()
    students = db.scalars(
        select(Student)
        .join(Enrollment, Enrollment.student_id == Student.id)
        .join(User, User.id == Student.user_id)
        .where(Enrollment.course_id == course.id)
        .order_by(User.full_name)
    ).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencia"
    header = (
        ["Registro", "Estudiante", "Email"]
        + [s.session_date.strftime("%d/%m") for s in sessions]
        + ["% Asistencia"]
    )
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
        c.alignment = Alignment(horizontal="center")
    by_session = {s.id: {r.student_id: r.status for r in s.attendance_records} for s in sessions}
    done = [s for s in sessions if is_closed(s)]
    done_ids = {s.id for s in done}
    for st in students:
        row = [st.user.registration_code, st.user.full_name, st.user.email or ""]
        attended = 0
        for s in sessions:
            status_ = by_session[s.id].get(st.id)
            if status_ is None and s.id not in done_ids:
                row.append("")
                continue
            status_ = status_ or AttendanceStatus.absent
            if status_ in (AttendanceStatus.present, AttendanceStatus.late):
                attended += 1
            row.append(STATUS_LABEL[status_])
        row.append(round(attended / len(done) * 100, 1) if done else 0)
        ws.append(row)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 32
    for i in range(4, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 8
    ws.freeze_panes = "D2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="asistencia_{course.code}.xlsx"'},
    )
