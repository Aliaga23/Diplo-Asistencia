"""Importación de estudiantes y horarios desde Excel (.xlsx)."""

import io
import re
from datetime import date, datetime, time
from uuid import UUID

from fastapi import HTTPException, status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.security import hash_password
from app.models import (
    ClassSession,
    Course,
    Enrollment,
    ImportBatch,
    ImportKind,
    Student,
    User,
    UserRole,
)

STUDENT_COLUMNS = {
    "registro": ["registro", "codigo", "código", "code", "registration_code", "matricula"],
    "nombre": ["nombre", "nombre completo", "name", "full_name", "estudiante"],
    "email": ["email", "correo", "correo electronico", "correo electrónico", "e-mail"],
    "carrera": ["carrera", "career", "programa", "program", "especialidad"],
}
SESSION_COLUMNS = {
    "fecha": ["fecha", "date", "dia", "día", "session_date"],
    "hora_inicio": ["hora_inicio", "hora inicio", "inicio", "start", "start_time", "hora"],
    "hora_fin": ["hora_fin", "hora fin", "fin", "end", "end_time"],
    "tema": ["tema", "topic", "contenido", "descripcion", "descripción"],
}

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _norm(v) -> str:
    return str(v).strip().lower() if v is not None else ""


def _map_headers(header_row, spec: dict[str, list[str]]) -> dict[str, int]:
    found: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        h = _norm(cell)
        for key, aliases in spec.items():
            if key not in found and h in aliases:
                found[key] = idx
    return found


def _load_rows(content: bytes):
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, "El archivo no es un Excel (.xlsx) válido"
        ) from exc
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c is not None and str(c).strip() != "" for c in r)]
    if not rows:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "El archivo está vacío")
    return rows[0], rows[1:]


def _row_reader(row: list, cols: dict[str, int]):
    def cell(key: str):
        idx = cols.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    return cell


def _parse_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"fecha inválida: {s}")


def _parse_time(v) -> time:
    if isinstance(v, datetime):
        return v.time().replace(second=0, microsecond=0)
    if isinstance(v, time):
        return v.replace(second=0, microsecond=0)
    if isinstance(v, int | float):  # fracción de día de Excel
        total = round(float(v) * 24 * 60)
        return time(hour=(total // 60) % 24, minute=total % 60)
    s = str(v).strip()
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%H.%M"):
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            continue
    raise ValueError(f"hora inválida: {s}")


def _get_course(db: Session, course_id: UUID) -> Course:
    course = db.get(Course, course_id)
    if course is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Curso no encontrado")
    return course


def import_students(
    db: Session, admin: User, course_id: UUID, filename: str, content: bytes
) -> ImportBatch:
    course = _get_course(db, course_id)
    header, rows = _load_rows(content)
    cols = _map_headers(header, STUDENT_COLUMNS)
    missing = [k for k in ("registro", "nombre") if k not in cols]
    if missing:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"Faltan columnas obligatorias: {', '.join(missing)}. "
            "Descarga la plantilla para ver el formato esperado.",
        )
    batch = ImportBatch(
        kind=ImportKind.students,
        filename=filename,
        course_id=course.id,
        uploaded_by=admin.id,
        rows_total=len(rows),
        rows_ok=0,
        rows_failed=0,
    )
    errors: list[dict] = []
    seen_emails: set[str] = set()
    seen_codes: set[str] = set()
    # Todas las cuentas nuevas reciben la contraseña común configurada (un solo hash por lote).
    default_hash = hash_password(get_settings().default_password)
    for i, row in enumerate(rows, start=2):
        cell = _row_reader(row, cols)

        code = str(cell("registro") or "").strip()
        if code.endswith(".0"):
            code = code[:-2]
        name = str(cell("nombre") or "").strip()
        email = _norm(cell("email")) or None
        career = cell("carrera")
        career = str(career).strip() if career not in (None, "") else None
        try:
            if not code or not name:
                raise ValueError("registro y nombre son obligatorios")
            if email and not EMAIL_RE.match(email):
                raise ValueError(f"email inválido: {email}")
            if code in seen_codes or (email and email in seen_emails):
                raise ValueError("fila duplicada dentro del archivo")
            seen_codes.add(code)
            if email:
                seen_emails.add(email)

            user = db.scalar(select(User).where(User.registration_code == code))
            if user is None:
                if email and db.scalar(select(User).where(User.email == email)):
                    raise ValueError(f"el email {email} ya pertenece a otro usuario")
                user = User(
                    registration_code=code,
                    email=email,
                    full_name=name,
                    role=UserRole.student,
                    password_hash=default_hash,
                )
                db.add(user)
                db.flush()
                student = Student(user_id=user.id, career=career)
                db.add(student)
                db.flush()
            else:
                if user.role != UserRole.student or user.student is None:
                    raise ValueError(f"el registro {code} pertenece a un usuario no estudiante")
                student = user.student
                if career and not student.career:
                    student.career = career
                if email and not user.email:
                    user.email = email
            exists = db.scalar(
                select(Enrollment).where(
                    Enrollment.student_id == student.id, Enrollment.course_id == course.id
                )
            )
            if exists is None:
                db.add(Enrollment(student_id=student.id, course_id=course.id))
            batch.rows_ok += 1
        except ValueError as exc:
            errors.append({"row": i, "message": str(exc)})
            batch.rows_failed += 1
    batch.errors = errors
    db.add(batch)
    db.commit()
    db.refresh(batch)
    return batch


def import_sessions(
    db: Session, admin: User, course_id: UUID, filename: str, content: bytes
) -> ImportBatch:
    course = _get_course(db, course_id)
    header, rows = _load_rows(content)
    cols = _map_headers(header, SESSION_COLUMNS)
    missing = [k for k in ("fecha", "hora_inicio", "hora_fin") if k not in cols]
    if missing:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"Faltan columnas obligatorias: {', '.join(missing)}. "
            "Descarga la plantilla para ver el formato esperado.",
        )
    batch = ImportBatch(
        kind=ImportKind.sessions,
        filename=filename,
        course_id=course.id,
        uploaded_by=admin.id,
        rows_total=len(rows),
        rows_ok=0,
        rows_failed=0,
    )
    errors: list[dict] = []
    for i, row in enumerate(rows, start=2):
        cell = _row_reader(row, cols)

        try:
            d = _parse_date(cell("fecha"))
            start = _parse_time(cell("hora_inicio"))
            end = _parse_time(cell("hora_fin"))
            if end <= start:
                raise ValueError("la hora de fin debe ser posterior a la de inicio")
            topic = cell("tema")
            topic = str(topic).strip() if topic not in (None, "") else None
            existing = db.scalar(
                select(ClassSession).where(
                    ClassSession.course_id == course.id,
                    ClassSession.session_date == d,
                    ClassSession.start_time == start,
                )
            )
            if existing is None:
                db.add(
                    ClassSession(
                        course_id=course.id,
                        session_date=d,
                        start_time=start,
                        end_time=end,
                        topic=topic,
                    )
                )
            else:
                existing.end_time = end
                if topic:
                    existing.topic = topic
            batch.rows_ok += 1
        except (ValueError, TypeError) as exc:
            errors.append({"row": i, "message": str(exc)})
            batch.rows_failed += 1
    batch.errors = errors
    db.add(batch)
    db.commit()
    db.refresh(batch)
    return batch


HEADER_FILL = PatternFill("solid", fgColor="1D4ED8")
EXAMPLE_FILL = PatternFill("solid", fgColor="EEF2FF")


def _model_workbook(
    title: str,
    headers: list[tuple[str, bool, str, int]],
    examples: list[list],
    instructions: list[str],
    number_formats: dict[int, str] | None = None,
) -> bytes:
    """Genera el modelo de importación: hoja de datos + hoja de instrucciones."""
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append([h[0] for h in headers])
    for idx, (_name, required, hint, width) in enumerate(headers, start=1):
        c = ws.cell(row=1, column=idx)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.comment = Comment(f"{'Obligatoria' if required else 'Opcional'}: {hint}", "Presencia")
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.row_dimensions[1].height = 22
    for row in examples:
        ws.append(row)
    for r in range(2, 2 + len(examples)):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=col)
            cell.fill = EXAMPLE_FILL
            if number_formats and col in number_formats:
                cell.number_format = number_formats[col]
    if number_formats:
        for col, fmt in number_formats.items():
            for r in range(2 + len(examples), 502):
                ws.cell(row=r, column=col).number_format = fmt
    if title == "horarios":
        dv = DataValidation(
            type="time",
            operator="between",
            formula1="0",
            formula2="0.999988",
            showErrorMessage=True,
            errorTitle="Hora inválida",
            error="Ingresa una hora en formato HH:MM (24 horas).",
        )
        ws.add_data_validation(dv)
        dv.add("B2:C500")
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Instrucciones")
    info.column_dimensions["A"].width = 100
    info.append([f"Modelo de importación de {title} — Presencia"])
    info["A1"].font = Font(bold=True, size=14, color="1D4ED8")
    info.append([""])
    for line in instructions:
        info.append([line])
    info.append([""])
    info.append(["Columnas:"])
    info.cell(row=info.max_row, column=1).font = Font(bold=True)
    for name, required, hint, _ in headers:
        info.append([f"  • {name} ({'obligatoria' if required else 'opcional'}): {hint}"])
    info.append([""])
    info.append(["Las filas sombreadas son ejemplos: reemplázalas o elimínalas antes de importar."])
    for row in info.iter_rows(min_row=3):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_students() -> bytes:
    return _model_workbook(
        "estudiantes",
        [
            ("registro", True, "número de registro único; es el usuario de ingreso", 16),
            ("nombre", True, "nombre completo del estudiante", 36),
            ("email", False, "correo institucional válido", 34),
            ("carrera", False, "carrera o especialidad del estudiante", 30),
        ],
        [
            [
                "218166141",
                "Ana Lucía Pérez Rojas",
                "ana.perez@ejemplo.edu.bo",
                "Ingeniería Informática",
            ],
            ["219045872", "Carlos Eduardo Vaca Suárez", "", "Ingeniería de Sistemas"],
        ],
        [
            "1. Completa una fila por estudiante en la hoja «Datos»,"
            " sin modificar los encabezados.",
            "2. Escribe el registro como texto (sin puntos ni espacios)."
            " Si ya existe, el estudiante solo se inscribe en el programa seleccionado.",
            "3. Todas las cuentas nuevas ingresan con su registro"
            " y la contraseña común del sistema.",
            "4. Guarda el archivo en formato .xlsx (máximo 5 MB) e impórtalo desde Presencia.",
        ],
        number_formats={1: "@"},
    )


def template_sessions() -> bytes:
    return _model_workbook(
        "horarios",
        [
            ("fecha", True, "fecha de la clase: AAAA-MM-DD o DD/MM/AAAA", 14),
            ("hora_inicio", True, "hora de inicio en formato HH:MM (24 h)", 13),
            ("hora_fin", True, "hora de fin HH:MM, posterior a la de inicio", 13),
            ("tema", False, "tema o contenido de la sesión", 44),
        ],
        [
            [date(2026, 10, 13), time(19, 0), time(21, 30), "Introducción al programa"],
            [date(2026, 10, 15), time(19, 0), time(21, 30), "Modelado de datos"],
        ],
        [
            "1. Completa una fila por sesión de clase en la hoja «Datos».",
            "2. Si una sesión con la misma fecha y hora de inicio ya existe en el programa, se"
            " actualizan su hora de fin y su tema.",
            "3. Los estudiantes solo pueden registrar su asistencia el mismo día de la sesión.",
            "4. Guarda el archivo en formato .xlsx (máximo 5 MB) e impórtalo desde Presencia.",
        ],
        number_formats={1: "dd/mm/yyyy", 2: "hh:mm", 3: "hh:mm"},
    )
