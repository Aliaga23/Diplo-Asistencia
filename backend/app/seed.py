"""Datos iniciales.

Uso:
    uv run python -m app.seed            # admin + programas + archivos de ejemplo en samples/
    uv run python -m app.seed --demo     # además importa los ejemplos y genera asistencia pasada

Todas las cuentas usan la contraseña común DEFAULT_PASSWORD (por defecto: admin1234!).
El administrador ingresa con el registro ADMIN_REGISTRATION (por defecto: 218166141).
"""

import argparse
import io
import random
import unicodedata
from dataclasses import dataclass
from datetime import date, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.security import hash_password
from app.core.timeutils import combine_local, today_local
from app.db.session import SessionLocal
from app.models import (
    AttendanceRecord,
    AttendanceSource,
    AttendanceStatus,
    ClassSession,
    Course,
    Enrollment,
    ImportBatch,
    Student,
    User,
    UserRole,
)
from app.services.imports import import_sessions, import_students

SAMPLES = Path(__file__).resolve().parent.parent / "samples"

STUDENTS = [
    ("219034521", "Ana Lucía Pérez Rojas", "Ingeniería Informática"),
    ("219045872", "Carlos Eduardo Vaca Suárez", "Ingeniería de Sistemas"),
    ("220011345", "María Fernanda Justiniano Áñez", "Ingeniería Informática"),
    ("218077123", "José Luis Mendoza Rivero", "Redes y Telecomunicaciones"),
    ("220034567", "Daniela Sofía Gutiérrez Ortiz", "Ingeniería de Sistemas"),
    ("219088234", "Luis Fernando Arteaga Melgar", "Ingeniería Informática"),
    ("221002345", "Valeria Alejandra Rocha Sandoval", "Ingeniería de Sistemas"),
    ("218066789", "Diego Armando Chávez Paz", "Redes y Telecomunicaciones"),
    ("220078901", "Camila Andrea Salvatierra Roca", "Ingeniería Informática"),
    ("219012678", "Jorge Andrés Cuéllar Moreno", "Ingeniería de Sistemas"),
    ("221045123", "Gabriela Nicole Antelo Flores", "Ingeniería Informática"),
    ("218023456", "Mauricio Javier Ribera Céspedes", "Ingeniería de Sistemas"),
    ("220056789", "Paola Andrea Hurtado Vargas", "Redes y Telecomunicaciones"),
    ("219067345", "Rodrigo Alonso Egüez Saucedo", "Ingeniería Informática"),
    ("221089012", "Natalia Belén Landívar Coimbra", "Ingeniería de Sistemas"),
    ("218045678", "Sebastián Ignacio Pinto Aguilera", "Ingeniería Informática"),
    ("220023891", "Mariana Isabel Callaú Toledo", "Ingeniería de Sistemas"),
    ("219056123", "Andrés Felipe Terrazas Lima", "Redes y Telecomunicaciones"),
    ("221067234", "Fabiola Cristina Montero Banegas", "Ingeniería Informática"),
    ("218089345", "Ricardo Manuel Dorado Guzmán", "Ingeniería de Sistemas"),
    ("220091122", "Lucía Esther Zambrana Rivas", "Ingeniería Informática"),
    ("219093344", "Fernando José Méndez Ortiz", "Ingeniería de Sistemas"),
    ("221095566", "Carla Patricia Soliz Aguirre", "Ingeniería Informática"),
    ("218097788", "Iván Gonzalo Peña Vaca", "Redes y Telecomunicaciones"),
    ("220099900", "Andrea Milena Rojas Suárez", "Ingeniería de Sistemas"),
    ("219081357", "Marco Antonio Gil Barba", "Ingeniería Informática"),
]

TOPICS_FS = [
    "Presentación del módulo y entorno de trabajo",
    "Modelado de datos y normalización",
    "Diseño de APIs REST con FastAPI",
    "Autenticación JWT y control de roles",
    "SQLAlchemy 2.0 y migraciones con Alembic",
    "Fundamentos de Next.js App Router",
    "Componentes con shadcn/ui y Tailwind",
    "Consumo de APIs y manejo de estado",
    "Importación y exportación de Excel",
    "Pruebas automatizadas con pytest",
    "Contenedores con Docker Compose",
    "Integración continua con GitHub Actions",
    "Despliegue y observabilidad",
    "Seguridad en aplicaciones web",
    "Revisión de proyectos finales",
    "Defensa de proyectos",
]

TOPICS_IA = [
    "Fundamentos de aprendizaje automático",
    "Preparación y limpieza de datos",
    "Regresión y clasificación supervisada",
    "Evaluación de modelos y métricas",
    "Redes neuronales y retropropagación",
    "Procesamiento de lenguaje natural",
    "Modelos de lenguaje y embeddings",
    "Visión por computadora",
    "Despliegue de modelos en producción",
    "Ética y gobernanza de la IA",
    "Proyecto integrador",
]


@dataclass(frozen=True)
class ProgramSeed:
    code: str
    name: str
    period: str
    students: tuple[int, ...]  # índices sobre STUDENTS
    weekdays: tuple[int, ...]  # 0 = lunes
    start: time
    end: time
    topics: list[str]
    sample_suffix: str


PROGRAMS = [
    ProgramSeed(
        code="DIP-FS-2026",
        name="Desarrollo Full Stack Empresarial",
        period="2026-II",
        students=tuple(range(20)),
        weekdays=(1, 3),
        start=time(19, 0),
        end=time(21, 30),
        topics=TOPICS_FS,
        sample_suffix="",
    ),
    ProgramSeed(
        code="DIP-IA-2026",
        name="Inteligencia Artificial Aplicada",
        period="2026-II",
        students=(0, 2, 4, 6, 8, 11, 13, 20, 21, 22, 23, 24, 25),
        weekdays=(1, 5),
        start=time(8, 0),
        end=time(10, 0),
        topics=TOPICS_IA,
        sample_suffix="_ia",
    ),
]


def _email(name: str) -> str:
    parts = name.split()
    first, last = parts[0], parts[-2] if len(parts) >= 3 else parts[-1]
    raw = unicodedata.normalize("NFKD", f"{first}.{last}".lower()).encode("ascii", "ignore")
    return f"{raw.decode()}@estudiantes.edu.bo"


def _style_header(ws) -> None:
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
        c.alignment = Alignment(horizontal="center")


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_students_xlsx(program: ProgramSeed) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(["registro", "nombre", "email", "carrera"])
    for i in program.students:
        code, name, career = STUDENTS[i]
        ws.append([code, name, _email(name), career])
    _style_header(ws)
    for col, w in zip("ABCD", (14, 38, 38, 30), strict=True):
        ws.column_dimensions[col].width = w
    return _save(wb)


def session_dates(program: ProgramSeed, today: date) -> list[date]:
    """Días de clase del programa: 6 semanas atrás y 2 adelante; incluye siempre HOY."""
    out: list[date] = []
    d = today - timedelta(weeks=6)
    while d <= today + timedelta(weeks=2):
        if d.weekday() in program.weekdays:
            out.append(d)
        d += timedelta(days=1)
    if today not in out:
        out.append(today)
    return sorted(out)


def build_sessions_xlsx(program: ProgramSeed, today: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(["fecha", "hora_inicio", "hora_fin", "tema"])
    for i, d in enumerate(session_dates(program, today)):
        ws.append([d, program.start, program.end, program.topics[i % len(program.topics)]])
    _style_header(ws)
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "dd/mm/yyyy"
        row[1].number_format = "hh:mm"
        row[2].number_format = "hh:mm"
    for col, w in zip("ABCD", (14, 12, 12, 48), strict=True):
        ws.column_dimensions[col].width = w
    return _save(wb)


def ensure_admin(db: Session) -> User:
    s = get_settings()
    admin = db.scalar(select(User).where(User.registration_code == s.admin_registration))
    if admin is None:
        admin = User(
            registration_code=s.admin_registration,
            email=s.admin_email,
            full_name=s.admin_name,
            role=UserRole.admin,
            password_hash=hash_password(s.default_password),
        )
        db.add(admin)
        db.commit()
        print(f"  ✓ admin creado: registro {s.admin_registration}")
    return admin


def ensure_course(db: Session, program: ProgramSeed) -> Course:
    course = db.scalar(select(Course).where(Course.code == program.code))
    if course is None:
        course = Course(code=program.code, name=program.name, period=program.period)
        db.add(course)
        db.commit()
        print(f"  ✓ programa creado: {course.code}")
    return course


def demo_attendance(db: Session, course: Course, admin: User, seed: int) -> int:
    """Genera asistencia histórica plausible para las sesiones ya dictadas."""
    rng = random.Random(seed)
    today = today_local()
    students = db.scalars(
        select(Student).join(Enrollment).where(Enrollment.course_id == course.id)
    ).all()
    profile = {st.id: rng.choice([0.97, 0.92, 0.88, 0.84, 0.78, 0.65]) for st in students}
    existing = {
        (r.session_id, r.student_id)
        for r in db.scalars(
            select(AttendanceRecord).join(ClassSession).where(ClassSession.course_id == course.id)
        ).all()
    }
    past = db.scalars(
        select(ClassSession).where(
            ClassSession.course_id == course.id, ClassSession.session_date < today
        )
    ).all()
    created = 0
    for s in past:
        for st in students:
            if (s.id, st.id) in existing:
                continue
            roll, p = rng.random(), profile[st.id]
            if roll < p * 0.84:
                status, minutes = AttendanceStatus.present, rng.randint(-10, 14)
            elif roll < p:
                status, minutes = AttendanceStatus.late, rng.randint(16, 45)
            elif roll < p + 0.05:
                status, minutes = AttendanceStatus.excused, 0
            else:
                continue  # falta: sin registro
            by_admin = status == AttendanceStatus.excused
            db.add(
                AttendanceRecord(
                    session_id=s.id,
                    student_id=st.id,
                    status=status,
                    source=AttendanceSource.admin if by_admin else AttendanceSource.self,
                    marked_at=combine_local(s.session_date, s.start_time)
                    + timedelta(minutes=minutes),
                    marked_by=admin.id if by_admin else st.user_id,
                    note="Justificación presentada" if by_admin else None,
                )
            )
            created += 1
    # Algunos estudiantes ya marcaron la sesión de hoy (actividad visible en el panel).
    today_session = db.scalar(
        select(ClassSession).where(
            ClassSession.course_id == course.id, ClassSession.session_date == today
        )
    )
    if today_session:
        for st in students[2:9]:
            if (today_session.id, st.id) in existing:
                continue
            db.add(
                AttendanceRecord(
                    session_id=today_session.id,
                    student_id=st.id,
                    status=AttendanceStatus.present,
                    source=AttendanceSource.self,
                    marked_at=combine_local(today, today_session.start_time)
                    - timedelta(minutes=rng.randint(1, 10)),
                    marked_by=st.user_id,
                )
            )
            created += 1
    db.commit()
    return created


def main(demo: bool) -> None:
    SAMPLES.mkdir(exist_ok=True)
    today = today_local()
    files: dict[str, tuple[bytes, bytes]] = {}
    for program in PROGRAMS:
        students_xlsx = build_students_xlsx(program)
        sessions_xlsx = build_sessions_xlsx(program, today)
        (SAMPLES / f"estudiantes{program.sample_suffix}.xlsx").write_bytes(students_xlsx)
        (SAMPLES / f"horarios{program.sample_suffix}.xlsx").write_bytes(sessions_xlsx)
        files[program.code] = (students_xlsx, sessions_xlsx)
    print(f"  ✓ archivos de ejemplo en {SAMPLES}")

    with SessionLocal() as db:
        admin = ensure_admin(db)
        for i, program in enumerate(PROGRAMS):
            course = ensure_course(db, program)
            if not demo:
                continue
            already = db.scalar(select(ImportBatch.id).where(ImportBatch.course_id == course.id))
            if already is not None:
                print(f"  · {course.code}: datos de demostración ya cargados, se omite")
                continue
            students_xlsx, sessions_xlsx = files[program.code]
            suffix = program.sample_suffix
            b1 = import_students(db, admin, course.id, f"estudiantes{suffix}.xlsx", students_xlsx)
            b2 = import_sessions(db, admin, course.id, f"horarios{suffix}.xlsx", sessions_xlsx)
            n = demo_attendance(db, course, admin, seed=42 + i)
            print(
                f"  ✓ {course.code}: {b1.rows_ok} estudiantes, {b2.rows_ok} sesiones,"
                f" {n} registros de asistencia"
            )
        if demo:
            print(f"  ✓ estudiante demo: registro {STUDENTS[0][0]} (inscrito en ambos programas)")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Carga datos iniciales")
    parser.add_argument("--demo", action="store_true", help="Importa datos y asistencia demo")
    main(parser.parse_args().demo)
